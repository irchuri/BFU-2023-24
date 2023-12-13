'''Дан файл ексель, в нем таблица с личными данными учеников разных школ (фамилия, имя, школа).
Перенести эту инфу в ворд, чтоб на каждом листе было "Награждается ученик такой-то школы по имени так-то" для каждой строки таблицы'''


import openpyxl
from docx import Document

EXCEL_FILENAME = 'Scolar_list.xlsx'
WORD_FILENAME = 'Scolar_gramoty.docx'


def excel_to_list(file_path):
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active

    data_list = []
    for row in sheet.iter_rows(values_only=True):
        data_list.append(list(row))

    workbook.close()
    return data_list


scolar_list = excel_to_list(EXCEL_FILENAME)

doc = Document()

for scolar in scolar_list:
    doc.add_paragraph(f'Награждается ученик {scolar[2]} школы {scolar[0]} {scolar[1]}')
    doc.add_page_break()
doc.save(WORD_FILENAME)
